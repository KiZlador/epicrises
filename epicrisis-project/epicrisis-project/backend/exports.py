import io
import json
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _flatten(result: dict) -> dict:
    """{группа: {признак: значение}} -> {(группа, признак): значение}"""
    flat = {}
    for key, value in result.items():
        if isinstance(value, dict):
            for sub_key, sub_value in value.items():
                flat[key, sub_key] = sub_value
        else:
            flat["", key] = value
    return flat


def _text(value) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def _set(ws, row: int, col: int, value) -> None:
    """Пишет значение строкой: ничего не превращается в число, дату или формулу."""
    cell = ws.cell(row=row, column=col, value=_text(value))
    cell.data_type = "s"


def build_json(result: dict) -> bytes:
    return json.dumps(result, ensure_ascii=False, indent=2).encode("utf-8")


def build_zip(data: dict) -> bytes:
    """Архив с отдельным JSON на каждый файл: result/<имя без .md>.json"""
    buf = io.BytesIO()
    used = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, result in data.items():
            stem, counter = (Path(name).stem, 1)
            unique = stem
            while unique in used:
                unique = f"{stem}_{counter}"
                counter += 1
            used.add(unique)
            zf.writestr(f"result/{unique}.json", build_json(result))
    return buf.getvalue()


def build_xlsx(data: dict) -> bytes:
    """Строки — файлы, столбцы — признаки; над названиями признаков — группы."""
    rows = {name: _flatten(result) for name, result in data.items()}
    columns: dict = {}
    for flat in rows.values():
        for col in flat:
            columns.setdefault(col, None)
    cols = list(columns)
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    _set(ws, 1, 1, "Файл")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    i = 0
    while i < len(cols):
        group, j = (cols[i][0], i)
        while j + 1 < len(cols) and cols[j + 1][0] == group:
            j += 1
        _set(ws, 1, 2 + i, group)
        if j > i:
            ws.merge_cells(start_row=1, start_column=2 + i, end_row=1, end_column=2 + j)
        i = j + 1
    for i, (_, key) in enumerate(cols):
        _set(ws, 2, 2 + i, key)
    for r, (name, flat) in enumerate(rows.items(), start=3):
        _set(ws, r, 1, name)
        for i, col in enumerate(cols):
            _set(ws, r, 2 + i, flat.get(col, ""))
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    for row in ws.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions["A"].width = min(40, max(len(n) for n in rows) + 2)
    for i, (_, key) in enumerate(cols):
        longest = max([len(key)] + [len(_text(f.get((_, key), ""))) for f in rows.values()])
        ws.column_dimensions[get_column_letter(2 + i)].width = min(45, longest + 2)
    ws.freeze_panes = "B3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
