import io
import json
import os
import subprocess
import sys
import tempfile
import textwrap
import unittest
import zipfile
from pathlib import Path
from threading import RLock
from types import SimpleNamespace

from openpyxl import load_workbook

from epicrises.documents import DocumentService, DocumentStore, ProcessingConflict
from epicrises.exporting import ExportService


class SessionCoordinator:
    def __init__(self):
        self.lock = RLock()
        self.active = False
        self.reset_sessions = []

    def busy(self, session_id):
        return self.active

    def reset(self, session_id):
        self.reset_sessions.append(session_id)


class ModuleBoundaries(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.storage = DocumentStore(Path(directory.name))
        self.coordinator = SessionCoordinator()
        self.documents = DocumentService("session", self.storage, self.coordinator)

    def upload(self, name):
        return self.documents.upload(
            [SimpleNamespace(filename=name, stream=io.BytesIO(b"Synthetic document"))]
        )

    def save(self, name, blood_pressure):
        result = {"осмотр при поступлении": {"bp": blood_pressure}}
        self.storage.save(name, SimpleNamespace(result=result, evidence={}, report={}))
        return result

    def test_document_snapshot_is_detached_from_later_updates_and_deletion(self):
        self.upload("ready.md")
        self.upload("pending.md")
        original = self.save("ready.md", "120/80")

        snapshot = self.documents.results_snapshot("all")
        self.assertEqual(snapshot, {"ready.md": original})
        self.assertEqual(self.documents.source("pending.md")["text"], "Synthetic document")
        self.assertIn("session", self.coordinator.reset_sessions)

        replacement = self.save("ready.md", "130/85")
        self.assertEqual(self.documents.result("ready.md"), replacement)
        self.assertEqual(snapshot, {"ready.md": original})

        self.coordinator.active = True
        with self.assertRaises(ProcessingConflict):
            self.documents.delete("ready.md")
        self.assertEqual(self.documents.result("ready.md"), replacement)

        self.coordinator.active = False
        self.documents.delete("ready.md")
        with self.assertRaises(FileNotFoundError):
            self.documents.result("ready.md")
        self.assertEqual(snapshot, {"ready.md": original})

    def test_export_needs_only_snapshot_even_after_documents_are_deleted(self):
        self.upload("ready.md")
        original = self.save("ready.md", "120/80")
        snapshot = self.documents.results_snapshot("one", "ready.md")
        self.documents.delete()
        self.assertEqual(self.documents.list(), [])

        exporter = ExportService()
        single = exporter.download(snapshot, "json", single=True)
        self.assertEqual(single.mime_type, "application/json")
        self.assertEqual(json.loads(single.content), original)

        archive = exporter.download(snapshot, "json")
        self.assertEqual(archive.mime_type, "application/zip")
        with zipfile.ZipFile(io.BytesIO(archive.content)) as files:
            self.assertEqual(json.loads(files.read("result/ready.json")), original)

        spreadsheet = exporter.download(snapshot, "xlsx")
        workbook = load_workbook(io.BytesIO(spreadsheet.content))
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.cell(3, 1).value, "ready.md")
        self.assertEqual(workbook.active.cell(3, 2).value, "120/80")
        self.assertEqual(snapshot, {"ready.md": original})

    def test_modules_import_without_web_or_model_runtime_from_another_directory(self):
        project_root = Path(__file__).resolve().parents[1]
        script = textwrap.dedent("""\
            import sys
            from pathlib import Path

            import epicrises.documents
            import epicrises.exporting
            assert "epicrises.processing" not in sys.modules

            import epicrises.processing
            from epicrises.shared.settings import ROOT, resolve_path

            assert "flask" not in sys.modules
            assert "llama_cpp" not in sys.modules
            assert ROOT == Path(sys.argv[1]).resolve()
            assert resolve_path("frontend/templates/index.html").is_file()
            assert resolve_path("models/example.gguf") == ROOT / "models/example.gguf"
            """)
        completed = subprocess.run(
            [sys.executable, "-c", script, str(project_root)],
            cwd=self.storage.root,
            env={**os.environ, "PYTHONPATH": str(project_root)},
            capture_output=True,
            text=True,
            timeout=15,
        )
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)


if __name__ == "__main__":
    unittest.main()
