import io
import json
import os
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from threading import Event
from unittest.mock import Mock, patch

from openpyxl import load_workbook

import pipeline as technical_cli
from backend.app import create_app
from backend.exports import build_xlsx
from domain.config import ROOT, Settings
from domain.schema import GROUPS
from domain.validation import validate_result
from extraction.evidence import browser_citations
from extraction.llm import LocalModel
from extraction.pipeline import Pipeline

DOC = """ЭПИКРИЗ
Поступил 01.01.2024. Выписан 04.01.2024.
Диагноз:
Сахарный диабет отрицает. Гипертоническая болезнь II стадии.
Первичный статус:
АД 119/83 мм рт. ст., пульс 77 в минуту. ЧДД 18.
Повторный осмотр:
АД 118/71 мм рт. ст., пульс 78 в минуту.
КАГ:
02.01.2024 выполнена селективная КАГ. ПМЖВ: стеноз 70%; ПКА: окклюзия.
Анализы крови:
Креатинин 99 мкмоль/л; лейкоциты 7,0 ×10⁹/л.
Контроль: креатинин 120 мкмоль/л.
"""


class ExtractionContracts(unittest.TestCase):
    def test_explicit_negation_initial_exam_and_first_lab(self):
        extraction = Pipeline(Settings(llm_enabled=False)).process(DOC, citations=True)
        result = extraction.result
        validate_result(result)
        self.assertEqual(sum(map(len, result.values())), 50)
        self.assertEqual(result["диагноз"]["dm"], "0")
        self.assertEqual(result["осмотр при поступлении"]["bp"], "119/83")
        self.assertEqual(result["осмотр при поступлении"]["bpm"], "77")
        self.assertEqual(result["Лабораторные данные"]["crea"], "99")
        self.assertEqual(result["коронарография"]["ca_date"], "02.01.2024")
        self.assertEqual(result["коронарография"]["ca_lad"], "1")
        self.assertEqual(result["коронарография"]["rca"], "2")

    def test_evidence_offsets_survive_unicode_and_units(self):
        raw = "🫀\r\n" + DOC.replace("7,0 ×10⁹/л", "7,0\u00a0×10⁹/л")
        extraction = Pipeline(Settings(llm_enabled=False)).process(raw, citations=True)
        citations = browser_citations(raw, extraction.evidence)
        self.assertIn("leucocytes", citations)
        encoded = raw.encode("utf-16-le")
        for group in extraction.evidence.values():
            for key, item in group.items():
                if item["verified"]:
                    self.assertEqual(raw[item["start"] : item["end"]], item["text"])
                    span = citations[key]
                    self.assertEqual(
                        encoded[2 * span["start"] : 2 * span["end"]].decode("utf-16-le"),
                        item["text"],
                    )
        self.assertFalse(extraction.evidence["диагноз"]["copd"]["verified"])

    def test_no_cag_means_no_vessel_codes(self):
        result = Pipeline(Settings(llm_enabled=False)).process("КАГ не выполнена.").result
        self.assertEqual(
            result["коронарография"],
            {"ca_fact": "N", "ca_date": "не указано", "ca_lad": "не указано", "rca": "не указано"},
        )

    def test_xlsx_keeps_text_and_does_not_execute_formulas(self):
        workbook = load_workbook(io.BytesIO(build_xlsx({"case.md": {"group": {"text": "=1+1"}}})))
        cell = workbook.active.cell(3, 2)
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s")

    def test_missing_model_does_not_silently_disable_llm(self):
        with self.assertRaises(RuntimeError):
            Pipeline(Settings(model_path="missing-model.gguf")).process(DOC)

    def test_llm_enforces_requested_keys(self):
        model = LocalModel(Settings())
        model._model = Mock()
        model._model.create_chat_completion.return_value = {
            "choices": [{"finish_reason": "stop", "message": {"content": '{"hf":"не указано"}'}}]
        }
        answer = model.review(DOC, ["hf"], {"hf": "не указано"}, {"triggered_rules": {}})
        self.assertEqual(answer, {"hf": "не указано"})
        schema = model._model.create_chat_completion.call_args.kwargs["response_format"]["schema"]
        self.assertEqual(schema["required"], ["hf"])
        self.assertFalse(schema["additionalProperties"])
        self.assertEqual(schema["properties"]["hf"]["type"], "string")

    def test_citation_failure_preserves_fields_and_reports_warning(self):
        pipeline = Pipeline(Settings(llm_enabled=False))
        pipeline.model = Mock()
        pipeline.model.review.return_value = {}
        pipeline.model.locate_evidence.side_effect = RuntimeError("citation service failed")
        result = pipeline.process(DOC, citations=True)
        validate_result(result.result)
        self.assertTrue(result.evidence["осмотр при поступлении"]["bp"]["verified"])
        self.assertTrue(any("citation service failed" in w for w in result.report["warnings"]))

    def test_model_cannot_put_statin_into_ace_field(self):
        pipeline = Pipeline(Settings(llm_enabled=False))
        pipeline.model = Mock()
        pipeline.model.review.return_value = {"ace_ing_sartan": "Аторвастатин 80 мг вечером"}
        result = pipeline.process(DOC)
        self.assertEqual(result.result["медикаментозная терапия"]["ace_ing_sartan"], "не указано")
        self.assertTrue(any("другого класса" in w for w in result.report["warnings"]))


class ApiContracts(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {"DATA_DIR": self.temp.name, "SECRET_KEY": "test-secret"})
        self.env.start()
        self.app = create_app(Settings(llm_enabled=False))
        self.app.testing = True
        self.client = self.app.test_client()

    def tearDown(self):
        self.env.stop()
        self.temp.cleanup()

    def upload(self, name="case.md"):
        return self.client.post("/upload", data={"files": (io.BytesIO(DOC.encode()), name)})

    def wait(self):
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            state = self.client.get("/process/status").get_json()
            if state["status"] not in ("queued", "running"):
                return state
            time.sleep(0.01)
        self.fail("Фоновая задача не завершилась")

    def test_upload_process_evidence_exports_delete_and_session_isolation(self):
        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"/static/files.js", page.data)
        self.assertEqual(self.upload().status_code, 200)
        self.assertEqual(self.app.test_client().get("/files").get_json()["files"], [])
        self.assertEqual(self.client.post("/process").status_code, 202)
        state = self.wait()
        self.assertEqual(state["succeeded"], 1, state)
        self.assertEqual(state["errors"], [])
        result = self.client.get("/result/case.md").get_json()
        validate_result(result)
        evidence = self.client.get("/evidence/case.md").get_json()
        source = self.client.get("/source/case.md").get_json()
        self.assertTrue(source["citations"])
        self.assertEqual(source["text"], DOC)
        for group, keys in GROUPS.items():
            for key in keys:
                self.assertEqual(result[group][key], evidence[group][key]["value"])
        self.assertEqual(
            self.client.get("/download?scope=one&format=json&name=case.md").get_json(), result
        )
        archive = zipfile.ZipFile(
            io.BytesIO(self.client.get("/download?scope=all&format=json").data)
        )
        self.assertEqual(archive.namelist(), ["result/case.json"])
        self.assertEqual(json.loads(archive.read("result/case.json")), result)
        workbook = load_workbook(
            io.BytesIO(self.client.get("/download?scope=all&format=xlsx").data)
        )
        self.assertEqual(workbook.active.max_column, 51)
        self.assertEqual(workbook.active.max_row, 3)
        self.assertEqual(self.client.delete("/files/case.md").status_code, 200)
        self.assertEqual(self.client.get("/files").get_json()["files"], [])
        self.assertEqual(self.client.get("/result/case.md").status_code, 404)
        self.assertFalse(list(Path(self.temp.name).rglob("*.json")))

    def test_zip_collision_and_invalid_utf8(self):
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as z:
            z.writestr("../case.md", DOC)
            z.writestr("folder/CASE.MD", DOC)
            z.writestr("readme.txt", "skip")
        response = self.client.post(
            "/upload", data={"files": (io.BytesIO(archive.getvalue()), "docs.zip")}
        )
        self.assertEqual(response.status_code, 200)
        names = response.get_json()["accepted"]
        self.assertEqual(len(names), 2)
        self.assertEqual(len({name.casefold() for name in names}), 2)
        self.assertFalse((Path(self.temp.name) / "case.md").exists())
        bad = self.client.post("/upload", data={"files": (io.BytesIO(b"\xff"), "bad.md")})
        self.assertEqual(bad.status_code, 400)
        self.assertEqual(len(self.client.get("/files").get_json()["files"]), 2)
        self.assertEqual(self.client.get("/download?format=unknown").status_code, 400)
        self.assertEqual(
            self.client.post(
                "/process", headers={"Origin": "https://elsewhere.invalid"}
            ).status_code,
            403,
        )

    def test_busy_job_blocks_mutations_and_does_not_duplicate(self):
        self.upload()
        entered, release = Event(), Event()
        original = Pipeline.process

        def delayed(pipeline, raw, **kwargs):
            entered.set()
            if not release.wait(5):
                raise RuntimeError("test timeout")
            return original(pipeline, raw, **kwargs)

        with patch.object(Pipeline, "process", delayed):
            try:
                started = self.client.post("/process").get_json()
                self.assertTrue(entered.wait(2))
                self.assertEqual(self.client.delete("/files").status_code, 409)
                self.assertEqual(self.upload("second.md").status_code, 409)
                self.assertEqual(self.client.post("/process").get_json()["id"], started["id"])
            finally:
                release.set()
                self.wait()

    def test_document_failure_does_not_abort_other_documents(self):
        self.upload("failed.md")
        self.upload("ok.md")
        original = Pipeline.process
        calls = 0

        def one_failure(pipeline, raw, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 1:
                raise ValueError("test document error")
            return original(pipeline, raw, **kwargs)

        with patch.object(Pipeline, "process", one_failure):
            self.client.post("/process")
            state = self.wait()
        self.assertEqual(state["succeeded"], 1)
        self.assertEqual(state["completed"], 2)
        self.assertEqual(len(state["errors"]), 1)
        self.assertEqual(self.client.get("/result/failed.md").status_code, 404)
        self.assertEqual(self.client.get("/result/ok.md").status_code, 200)


class CliContracts(unittest.TestCase):
    def test_interruption_preserves_report(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docs = root / "docs"
            docs.mkdir()
            (docs / "one.md").write_text(DOC)
            metrics = root / "metrics.json"
            args = [
                "pipeline.py",
                "--regex-only",
                "--input",
                str(docs),
                "--output",
                str(root / "results"),
                "--metrics",
                str(metrics),
            ]
            with (
                patch.object(sys, "argv", args),
                patch.object(Pipeline, "process", side_effect=KeyboardInterrupt),
            ):
                self.assertEqual(technical_cli.main(), 130)
            report = json.loads(metrics.read_text())
            self.assertTrue(report["interrupted"])
            self.assertEqual(report["succeeded"], 0)
            self.assertEqual(report["failed"], 0)
            self.assertEqual(report["records"][0]["status"], "interrupted")

    def test_batch_outputs_only_fields_and_failure_removes_stale_results(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docs, out = root / "docs", root / "results"
            docs.mkdir()
            out.mkdir()
            (docs / "one.md").write_text(DOC)
            (docs / "bad.md").write_bytes(b"\xff")
            (out / "bad.json").write_text("{}")
            (out / "one.ev.json").write_text("{}")
            metrics = root / "metrics.json"
            run = subprocess.run(
                [
                    sys.executable,
                    str(ROOT / "pipeline.py"),
                    "--regex-only",
                    "--input",
                    str(docs),
                    "--output",
                    str(out),
                    "--metrics",
                    str(metrics),
                ],
                capture_output=True,
                text=True,
                timeout=30,
            )
            self.assertEqual(run.returncode, 1, run.stderr)
            self.assertEqual([p.name for p in out.iterdir()], ["one.json"])
            validate_result(json.loads((out / "one.json").read_text()))
            report = json.loads(metrics.read_text())
            self.assertEqual(
                (report["documents"], report["succeeded"], report["failed"]), (2, 1, 1)
            )
            self.assertGreater(report["total_seconds"], 0)
            self.assertAlmostEqual(
                report["average_seconds_per_document"], report["total_seconds"] / 2
            )


if __name__ == "__main__":
    unittest.main()
